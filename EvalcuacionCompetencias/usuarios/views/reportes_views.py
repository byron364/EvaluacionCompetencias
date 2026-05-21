from django.http import HttpResponse, JsonResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from django.db.models import Count
from ..models import Perfil, Curso, Calificacion


def exportar_reporte(request, tipo):
    wb = Workbook()
    ws = wb.active
    ws.title = tipo.capitalize()

    if tipo == "usuarios":
        ws.append(["Nombre", "Apellido", "Correo", "Rol", "Documento", "Grado", "Unidad", "Activo"])
        for perfil in Perfil.objects.select_related("user").order_by("rol", "user__last_name"):
            ws.append([
                perfil.user.first_name,
                perfil.user.last_name,
                perfil.user.email,
                perfil.rol,
                perfil.documento,
                perfil.grado or "",
                perfil.unidad or "",
                "Sí" if perfil.user.is_active else "No",
            ])
    elif tipo == "cursos":
        ws.append(["Código", "Curso", "Instructor", "Inicio", "Fin", "Cupo", "Inscritos", "Activo"])
        cursos = Curso.objects.select_related("instructor").annotate(total_soldados=Count("inscripciones", distinct=True))
        for curso in cursos:
            ws.append([
                curso.codigo,
                curso.nombre,
                curso.instructor.get_full_name() if curso.instructor else "Sin instructor",
                curso.fecha_inicio,
                curso.fecha_fin,
                curso.cupo_maximo,
                curso.total_soldados,
                "Sí" if curso.activo else "No",
            ])
    elif tipo == "resultados":
        ws.append(["Soldado", "Documento", "Curso", "Actividad", "Nota", "Observaciones", "Fecha"])
        calificaciones = Calificacion.objects.select_related("estudiante", "estudiante__perfil", "curso", "tarea", "evaluacion")
        for calificacion in calificaciones:
            actividad = calificacion.evaluacion.titulo if calificacion.evaluacion else calificacion.tarea.titulo if calificacion.tarea else "General"
            ws.append([
                calificacion.estudiante.get_full_name(),
                calificacion.estudiante.perfil.documento,
                calificacion.curso.nombre,
                actividad,
                float(calificacion.nota),
                calificacion.observaciones or "",
                calificacion.creado_en.strftime("%Y-%m-%d"),
            ])
    else:
        return JsonResponse({"error": "Reporte no válido"}, status=404)

    for row in ws.iter_rows(min_row=1, max_row=1):
        for cell in row:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="789441")
            cell.alignment = Alignment(horizontal="center")

    for column_cells in ws.columns:
        ws.column_dimensions[column_cells[0].column_letter].width = 22

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="reporte_{tipo}.xlsx"'
    wb.save(response)
    return response