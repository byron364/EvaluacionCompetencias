import json
import base64
import os
import subprocess
import pandas as pd

from django.shortcuts import render
from django.http import JsonResponse, HttpResponse
from django.views.decorators.csrf import csrf_exempt
from django.contrib.auth.models import User
from django.contrib.auth.hashers import make_password
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image
from django.conf import settings
from ..models import *
from ..models import Perfil
from ..models import Batallon
from ..models import Compania
from django.core.files.base import ContentFile


def admin_usuarios(request):

    usuarios = Perfil.objects.select_related(
        'user'
    ).all()

    companias = Compania.objects.filter(
    activa=True
    )

    return render(

        request,

        "admin_usuarios.html",

        {

            "usuarios": usuarios,

            "companias": companias

        }

    )

def listar_usuarios(request):

    usuarios = Perfil.objects.select_related(
        "user"
    ).all()

    return render(
        request,
        "admin_usuarios.html",
        {
            "usuarios": usuarios
        }
    )


@csrf_exempt
def crear_usuario(request):

    if request.method != 'POST':

        return JsonResponse({
            'error': 'Método no permitido'
        }, status=405)

    try:

        # =========================
        # LEER DATA
        # =========================

        data = json.loads(request.body)

        nombres = data.get('nombres')
        apellidos = data.get('apellidos')
        email = data.get('email')
        documento = data.get('documento')
        grado = data.get('grado')
        rol = data.get('rol')
        compania_id = data.get("compania")

        # =========================
        # VALIDAR COMPAÑÍA
        # =========================

        compania = None

        if compania_id:

            compania = Compania.objects.filter(
                id=compania_id
            ).first()

        # =========================
        # VALIDAR CAMPOS
        # =========================

        campos_obligatorios = [

            nombres,
            apellidos,
            email,
            documento,
            rol,
            compania_id

        ]

        if rol != "soldado":

            campos_obligatorios.append(
                grado
            )

        if not all(campos_obligatorios):

            return JsonResponse({

                'error':
                    'Todos los campos son obligatorios'

            }, status=400)

        # =========================
        # VALIDAR CORREO
        # =========================

        if rol == "soldado":

            if not email.endswith(
                "@buzonejercito.mil.co"
            ):

                return JsonResponse({

                    "error":
                        "El soldado debe usar correo @buzonejercito.mil.co"

                }, status=400)

        else:

            if not email.endswith(
                "@ejercito.mil.co"
            ):

                return JsonResponse({

                    "error":
                        "Debe usar correo institucional @ejercito.mil.co"

                }, status=400)

        # =========================
        # VALIDAR USER
        # =========================

        if User.objects.filter(
            username__iexact=email
        ).exists():

            return JsonResponse({

                'error':
                    'El usuario ya existe'

            }, status=400)

        # =========================
        # VALIDAR DOCUMENTO
        # =========================

        if Perfil.objects.filter(
            documento__iexact=documento
        ).exists():

            return JsonResponse({

                'error':
                    'Documento ya registrado'

            }, status=400)

        # =========================
        # VALIDAR HUELLA TEMPORAL
        # =========================

        ruta_temp = rf"C:\Huellero\temp\temp_huella_{documento}.json"

        if not os.path.exists(ruta_temp):

            return JsonResponse({

                'error':
                    'Debe registrar la huella antes de crear el usuario'

            }, status=400)

        # =========================
        # CREAR USUARIO
        # =========================

        user = User.objects.create(

            username=email,
            email=email,

            first_name=nombres,
            last_name=apellidos,

            password=make_password(
                documento
            )
        )

        # =========================
        # CREAR PERFIL
        # =========================

        perfil = Perfil.objects.create(

            user=user,

            rol=rol,

            documento=documento,

            grado=grado,

            compania=compania,

        )

        # =========================
        # LEER HUELLA TEMPORAL
        # =========================

        with open(
            ruta_temp,
            "r",
            encoding="utf-8"
        ) as f:

            datos_huella = json.load(f)

        # =========================
        # TEMPLATE
        # =========================

        perfil.huella = datos_huella.get(
            "huella"
        )

        # =========================
        # IMAGEN BASE64
        # =========================

        imagen_base64 = datos_huella.get(
            "imagen"
        )

        # =========================
        # GUARDAR IMAGEN
        # =========================

        if imagen_base64:

            formato, imgstr = (

                imagen_base64.split(';base64,')

                if ';base64,' in imagen_base64

                else ('', imagen_base64)

            )

            image_data = ContentFile(

                base64.b64decode(
                    imgstr
                ),

                name=f"huella_{documento}.png"

            )

            perfil.imagen_huella = image_data

        # =========================
        # GUARDAR PERFIL
        # =========================

        perfil.save()

        # =========================
        # ELIMINAR JSON TEMPORAL
        # =========================

        if os.path.exists(ruta_temp):

            os.remove(ruta_temp)

        # =========================
        # RESPUESTA
        # =========================

        return JsonResponse({

            'success': True,

            'mensaje':
                'Usuario creado correctamente con huella biométrica'

        })

    except Exception as e:

        return JsonResponse({

            'error': str(e)

        }, status=500)

@csrf_exempt
def eliminar_usuario(request, user_id):

    try:

        user = User.objects.get(id=user_id)

        user.delete()

        return JsonResponse({
            "mensaje":
                "Usuario eliminado correctamente"
        })

    except User.DoesNotExist:

        return JsonResponse({
            "error":
                "Usuario no encontrado"
        }, status=404)


@csrf_exempt
def editar_usuario(request, user_id):

    if request.method == "POST":

        try:

            user = User.objects.get(id=user_id)

            perfil = Perfil.objects.get(user=user)

            data = json.loads(request.body)

            user.first_name = data.get("nombre")
            user.last_name = data.get("apellido")
            user.email = data.get("email")
            user.is_active = data.get("activo")

            perfil.rol = data.get("rol")

            user.save()

            perfil.save()

            return JsonResponse({
                "mensaje":
                    "Usuario actualizado correctamente"
            })

        except User.DoesNotExist:

            return JsonResponse({
                "error":
                    "Usuario no encontrado"
            }, status=404)

@csrf_exempt
def cargar_usuarios_excel(request):

    if request.method != "POST":

        return JsonResponse({
            "error": "Método no permitido"
        }, status=405)

    archivo = request.FILES.get("archivo")

    compania_id = request.POST.get(
        "compania"
    )

    if not compania_id:

        return JsonResponse({
            "error": "Debe seleccionar una compañía"
        }, status=400)

    compania = Compania.objects.filter(
        id=compania_id
    ).first()

    if not compania:

        return JsonResponse({
            "error": "Compañía no encontrada"
        }, status=400)

    batallon = compania.batallon

    if not archivo:

        return JsonResponse({
            "error": "No se recibió ningún archivo"
        }, status=400)

    try:

        df = pd.read_excel(
            archivo,
            header=7
        )

        creados = 0

        errores = []

        for index, row in df.iterrows():

            nombres = str(
                row.get("nombres", "")
            ).strip()

            apellidos = str(
                row.get("apellidos", "")
            ).strip()

            email = str(
                row.get("email", "")
            ).strip().lower()

            email = email.replace(
                " ",
                ""
            )

            documento = str(
                row.get("documento", "")
            ).replace(".0", "").strip()

            documento = documento.replace(
                " ",
                ""
            )

            grado = str(
                row.get("grado", "")
            ).strip()

            # =====================================
            # VALIDAR CAMPOS
            # =====================================

            if not nombres:

                errores.append(
                    f"Fila {index+9}: nombres vacíos"
                )

                continue

            if not apellidos:

                errores.append(
                    f"Fila {index+9}: apellidos vacíos"
                )

                continue

            if not email:

                errores.append(
                    f"Fila {index+9}: email vacío"
                )

                continue

            if not documento:

                errores.append(
                    f"Fila {index+9}: documento vacío"
                )

                continue

            # =====================================
            # VALIDAR CORREOS Y ROL
            # =====================================

            if email.endswith(
                "@buzonejercito.mil.co"
            ):

                rol = "soldado"

            elif email.endswith(
                "@ejercito.mil.co"
            ):

                rol = "instructor"

            else:

                errores.append(
                    f"Fila {index+9}: correo no permitido"
                )

                continue

            # =====================================
            # VALIDAR DUPLICADOS
            # =====================================

            if User.objects.filter(
                username=email
            ).exists():

                errores.append(
                    f"Fila {index+9}: usuario ya existe"
                )

                continue

            if Perfil.objects.filter(
                documento=documento
            ).exists():

                errores.append(
                    f"Fila {index+9}: documento ya registrado"
                )

                continue

            # =====================================
            # CREAR USUARIO
            # =====================================

            user = User.objects.create(

                username=email,

                email=email,

                first_name=nombres,

                last_name=apellidos,

                password=make_password(
                    documento
                )
            )

            # =====================================
            # CREAR PERFIL
            # =====================================

            Perfil.objects.create(

                user=user,

                rol=rol,

                documento=documento,

                grado=grado,

                compania=compania,

                batallon=batallon
            )

            creados += 1

        return JsonResponse({

            "mensaje":
                f"{creados} usuarios creados",

            "errores":
                errores
        })

    except Exception as e:

        import traceback

        traceback.print_exc()

        return JsonResponse({

            "error": str(e)

        }, status=500)
       
@csrf_exempt
def descargar_plantilla_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Plantilla Usuarios"

    ruta_logo_principal = os.path.join(
        settings.BASE_DIR,
        "static",
        "img",
        "escudo.png"
    )

    if os.path.exists(ruta_logo_principal):
        logo_principal = Image(ruta_logo_principal)
        logo_principal.width = 140
        logo_principal.height = 140

        # Logo izquierdo superior
        ws.add_image(logo_principal, "A1")


    ruta_logo_secundario = os.path.join(
        settings.BASE_DIR,
        "static",
        "img",
        "CompromisoColombia.png"
    )

    if os.path.exists(ruta_logo_secundario):
        logo_secundario = Image(ruta_logo_secundario)
        logo_secundario.width = 180
        logo_secundario.height = 110

        # Logo derecho superior
        ws.add_image(logo_secundario, "E2")

    encabezados = [
        "nombres",
        "apellidos",
        "email",
        "documento",
        "grado"
    ]

    for i, titulo in enumerate(encabezados, start=1):
        ws.cell(row=8, column=i, value=titulo)


    ejemplo = [
        "Juan",
        "Perez",
        "juan@ejercito.mil.co",
        "12345678",
        "Capitán"
    ]

    ws["A9"] = "Juan"
    ws["B9"] = "Perez"
    ws["C9"] = "juan@buzonejercito.mil.co"
    ws["D9"] = "12345678"
    ws["E9"] = "Capitán"

    for i, valor in enumerate(ejemplo, start=1):
        ws.cell(row=9, column=i, value=valor)

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="789441"
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
        size=12
    )

    center = Alignment(
        horizontal="center",
        vertical="center"
    )

    thin = Side(
        style="thin",
        color="D9D9D9"
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin
    )

    for col in ws[8]:
        col.fill = header_fill
        col.font = header_font
        col.alignment = center
        col.border = border


    for row in ws.iter_rows(min_row=9, max_row=9):
        for cell in row:
            cell.border = border
            cell.alignment = center

    widths = {
        "A": 22,
        "B": 22,
        "C": 35,
        "D": 20,
        "E": 28
    }

    for col, width in widths.items():
        ws.column_dimensions[col].width = width


    for fila in range(10, 501):
        for col in ["A", "B", "C", "D", "E"]:
            ws[f"{col}{fila}"] = ""

    ws.protection.sheet = True
    ws.protection.enable()

    # desbloquear celdas editables
    for fila in range(9, 501):
        for col in ["A", "B", "C", "D", "E"]:
            ws[f"{col}{fila}"].protection = (
                ws[f"{col}{fila}"].protection.copy(locked=False)
            )

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response["Content-Disposition"] = (
        'attachment; filename="plantilla_usuarios.xlsx"'
    )

    wb.save(response)

    return response


@csrf_exempt
def guardar_huella(request):

    if request.method == "POST":

        data = json.loads(
            request.body
        )

        documento = data.get(
            "documento"
        )

        huella = data.get(
            "huella"
        )

        imagen = data.get(
            "imagen"
        )

        try:

            perfil = Perfil.objects.get(
                documento=documento
            )

            # GUARDAR TEMPLATE
            perfil.huella = huella

            # GUARDAR IMAGEN
            image_data = ContentFile(
                base64.b64decode(imagen),
                name=f"huella_{perfil.documento}.png"
            )

            perfil.imagen_huella = image_data

            perfil.save()

            return JsonResponse({
                "success": True
            })

        except Perfil.DoesNotExist:

            return JsonResponse({
                "success": False,
                "error": "Usuario no encontrado"
            })

    return JsonResponse({
        "success": False
    })

def abrir_huellero(request, documento):

    ruta_exe = r"C:\Ejercito\EvalcuacionCompetencias\EvalcuacionCompetencias\Huellero\HuelleroMilitarRegistro.exe"

    subprocess.Popen([
        ruta_exe,
        documento
    ])

    return JsonResponse({
        "success": True
    })

def validar_huella_temp(request, documento):

    ruta_temp = rf"C:\Huellero\temp\temp_huella_{documento}.json"

    if not os.path.exists(ruta_temp):

        return JsonResponse({
            "success": False
        })

    try:

        with open(
            ruta_temp,
            "r",
            encoding="utf-8"
        ) as f:

            datos = json.load(f)

        return JsonResponse({

            "success": True,

            "imagen":
                datos.get("imagen")

        })

    except Exception as e:

        return JsonResponse({

            "success": False,

            "error": str(e)

        }, status=500)