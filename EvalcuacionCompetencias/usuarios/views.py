from django.shortcuts import render, redirect
from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.contrib.auth.models import User
from django.contrib.auth.hashers import make_password
from django.contrib.auth import authenticate, login, logout
from django.contrib import messages
from django.http import HttpResponse 
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from django.contrib.auth.decorators import login_required
from openpyxl.drawing.image import Image
import os
from django.conf import settings 
from .models import Perfil, ConfiguracionUsuario
from .helpers import obtener_configuracion_usuario
import json
import pandas as pd 
from .models import Curso, Inscripcion, Evaluacion, Calificacion
from django.db.models import Count, Avg, Q


def es_peticion_spa(request):
    return request.headers.get("X-Requested-With") == "XMLHttpRequest"


def contexto_admin():
    return {
        "instructores": User.objects.filter(perfil__rol="instructor")
    }


def login_view(request):
    if request.method == 'POST':

        email = (request.POST.get('email') or '').strip().lower()
        password = request.POST.get('password') or ''

        if not email.endswith('@ejercito.mil.co'):
            messages.error(request, 'Correo no válido')
            return render(request, 'layout/login.html')

        user = authenticate(request, username=email, password=password)

        if user is None:
            try:
                usuario = User.objects.get(email__iexact=email)
                user = authenticate(request, username=usuario.username, password=password)
            except User.DoesNotExist:
                user = None

        if user is not None:
            login(request, user)

            try:
                rol = user.perfil.rol
            except:
                messages.error(request, 'Usuario sin rol')
                return redirect('login')

            if rol == 'admin':
                return redirect('admin_dashboard')
            elif rol == 'instructor':
                return redirect('instructor_dashboard')
            elif rol == 'soldado':
                return redirect('soldado_dashboard')

        else:
            messages.error(request, 'Credenciales incorrectas')

    return render(request, 'layout/login.html')



def registro_view(request):
    if request.method == 'POST':

        email = (request.POST.get('email') or '').strip().lower()
        password = request.POST.get('password') or ''
        rol = request.POST.get('rol')

        if not email.endswith('@ejercito.mil.co'):
            return render(request, 'registro.html', {
                'error': 'Debe usar correo institucional'
            })

        user = User.objects.create(
            username=email,
            email=email,
            first_name=nombres, 
            last_name=apellidos,
            password=make_password(password)
        )

        Perfil.objects.create(user=user, rol=rol)

        return redirect('login')

    return render(request, 'registro.html')


def rol_requerido(rol_permitido):
    def decorator(view_func):
        def wrapper(request, *args, **kwargs):
            if not request.user.is_authenticated:
                return redirect('login')

            if request.user.perfil.rol != rol_permitido:
                return redirect('login')

            return view_func(request, *args, **kwargs)
        return wrapper
    return decorator



@rol_requerido('admin')
def admin_dashboard(request):
    return render(request, "admin.html", {
        **contexto_admin(),
        "configuracion": obtener_configuracion_usuario(request.user)
    })

@rol_requerido('instructor')
def instructor_dashboard(request):
    return render(request, 'instructor.html', {
        "configuracion": obtener_configuracion_usuario(request.user)
    })


@rol_requerido('instructor')
def instructor_inicio(request):
    if not es_peticion_spa(request):
        return render(request, 'instructor.html', {
            "configuracion": obtener_configuracion_usuario(request.user)
        })

    cursos = Curso.objects.filter(instructor=request.user)
    curso_ids = cursos.values_list("id", flat=True)
    soldados_asignados = User.objects.filter(mis_cursos__curso__in=curso_ids).distinct().count()
    evaluaciones_pendientes = Evaluacion.objects.filter(curso__in=cursos).count()
    promedio = Calificacion.objects.filter(curso__in=cursos).aggregate(promedio=Avg("nota"))["promedio"]

    return render(request, 'instructor_inicio.html', {
        "configuracion": obtener_configuracion_usuario(request.user),
        "total_cursos": cursos.filter(activo=True).count(),
        "soldados_asignados": soldados_asignados,
        "evaluaciones_pendientes": evaluaciones_pendientes,
        "promedio_general": round(float(promedio), 1) if promedio is not None else 0,
        "cursos": cursos.annotate(total_soldados=Count("inscripciones", distinct=True))[:5],
        "calificaciones": Calificacion.objects.filter(curso__in=cursos).select_related("estudiante", "curso").order_by("-creado_en")[:5],
    })


@rol_requerido('instructor')
def instructor_cursos(request):
    if not es_peticion_spa(request):
        return render(request, 'instructor.html')

    cursos = Curso.objects.filter(instructor=request.user).annotate(
        total_soldados=Count("inscripciones", distinct=True)
    ).order_by("-creado_en")
    return render(request, 'instructor_cursos.html', {"cursos": cursos})


@rol_requerido('instructor')
def instructor_evaluaciones(request):
    if not es_peticion_spa(request):
        return render(request, 'instructor.html')

    evaluaciones = Evaluacion.objects.filter(curso__instructor=request.user).select_related("curso").order_by("-fecha")
    return render(request, 'instructor_evaluaciones.html', {"evaluaciones": evaluaciones})


@rol_requerido('instructor')
def instructor_soldados(request):
    if not es_peticion_spa(request):
        return render(request, 'instructor.html')

    inscripciones = Inscripcion.objects.filter(curso__instructor=request.user).select_related(
        "curso", "estudiante", "estudiante__perfil"
    ).order_by("estudiante__last_name", "estudiante__first_name")
    return render(request, 'instructor_soldados.html', {"inscripciones": inscripciones})


@rol_requerido('instructor')
def instructor_reportes(request):
    if not es_peticion_spa(request):
        return render(request, 'instructor.html')

    cursos = Curso.objects.filter(instructor=request.user).annotate(
        total_soldados=Count("inscripciones", distinct=True),
        promedio=Avg("calificacion__nota")
    )
    return render(request, 'instructor_reportes.html', {"cursos": cursos})


@rol_requerido('soldado')
def soldado_dashboard(request):
    return render(request, 'soldado.html')



def logout_view(request):
    logout(request)
    return redirect('login')

def cursos(request):
    if not request.user.is_authenticated:
        return redirect('login')

    rol = getattr(request.user.perfil, "rol", None)

    if rol == "admin":
        if not es_peticion_spa(request):
            return render(request, "admin.html", contexto_admin())
        return admin_cursos(request)

    if rol == "soldado":
        if not es_peticion_spa(request):
            return render(request, 'soldado.html')
        inscripciones = Inscripcion.objects.filter(estudiante=request.user).select_related("curso", "curso__instructor")
        return render(request, 'cursos.html', {"inscripciones": inscripciones})

    if rol == "instructor":
        if not es_peticion_spa(request):
            return render(request, 'instructor.html')
        return instructor_cursos(request)

    return redirect('login')

def evaluaciones(request):
    if not request.user.is_authenticated:
        return redirect('login')
    if not es_peticion_spa(request):
        return render(request, 'soldado.html')

    evaluaciones_usuario = Evaluacion.objects.filter(curso__inscripciones__estudiante=request.user).select_related("curso").distinct()
    return render(request, 'evaluaciones.html', {"evaluaciones": evaluaciones_usuario})

def resultados(request):
    if not request.user.is_authenticated:
        return redirect('login')
    if not es_peticion_spa(request):
        return render(request, 'soldado.html')

    calificaciones = Calificacion.objects.filter(estudiante=request.user).select_related("curso", "tarea", "evaluacion").order_by("-creado_en")
    promedio = calificaciones.aggregate(promedio=Avg("nota"))["promedio"]
    return render(request, 'resultados.html', {
        "calificaciones": calificaciones,
        "promedio": round(float(promedio), 1) if promedio is not None else 0
    })

def retroalimentacion(request):
    if request.user.is_authenticated and not es_peticion_spa(request):
        return render(request, 'soldado.html')

    return render(request, 'retro.html')

def inicio(request):
    if not request.user.is_authenticated:
        return redirect('login')
    if not es_peticion_spa(request):
        return render(request, 'soldado.html')

    inscripciones = Inscripcion.objects.filter(estudiante=request.user)
    calificaciones = Calificacion.objects.filter(estudiante=request.user)
    promedio = calificaciones.aggregate(promedio=Avg("nota"))["promedio"]
    return render(request, 'inicio.html', {
        "total_cursos": inscripciones.count(),
        "total_evaluaciones": Evaluacion.objects.filter(curso__inscripciones__estudiante=request.user).distinct().count(),
        "total_resultados": calificaciones.count(),
        "promedio": round(float(promedio), 1) if promedio is not None else 0,
    })


def admin_usuarios(request):
    if request.user.is_authenticated and not es_peticion_spa(request):
        return render(request, "admin.html", contexto_admin())

    usuarios = Perfil.objects.select_related("user").all()

    return render(
        request,
        "admin_usuarios.html",
        {
            "usuarios": usuarios
        }
    )


def admin_cursos(request):
    if request.user.is_authenticated and not es_peticion_spa(request):
        return render(request, "admin.html", contexto_admin())

    if request.user.is_authenticated and getattr(request.user.perfil, "rol", None) == "soldado":
        return cursos(request)

    if request.user.is_authenticated and getattr(request.user.perfil, "rol", None) == "instructor":
        return instructor_cursos(request)

    instructores = User.objects.filter(
        perfil__rol="instructor"
    )

    soldados = User.objects.filter(
        perfil__rol="soldado",
        is_active=True
    ).select_related("perfil").order_by("last_name", "first_name")

    cursos = Curso.objects.select_related(
        'instructor'
    ).annotate(total_soldados=Count("inscripciones", distinct=True)).all()

    return render(
        request,
        "admin_cursos.html",
        {
            "instructores": instructores,
            "cursos": cursos,
            "soldados": soldados
        }
    )

def admin_reportes(request):
    if request.user.is_authenticated and not es_peticion_spa(request):
        return render(request, "admin.html", contexto_admin())

    resumen = {
        "usuarios": User.objects.count(),
        "soldados": Perfil.objects.filter(rol="soldado").count(),
        "instructores": Perfil.objects.filter(rol="instructor").count(),
        "cursos": Curso.objects.count(),
        "inscripciones": Inscripcion.objects.count(),
        "promedio": Calificacion.objects.aggregate(promedio=Avg("nota"))["promedio"] or 0,
    }
    return render(request, 'admin_reportes.html', {"resumen": resumen})
def admin_dashboard_partial(request):
    return render(request, 'admin_dashboard_partial.html', {
        "configuracion": obtener_configuracion_usuario(request.user)
    })

@csrf_exempt
def crear_usuario(request):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)

            nombres = data.get('nombres')
            apellidos = data.get('apellidos')
            email = data.get('email')
            documento = data.get('documento')
            unidad = data.get('unidad')
            grado = data.get('grado')
            rol = data.get('rol')

            
            if not all([nombres, apellidos, email, documento, rol]):
                return JsonResponse({'error': 'Todos los campos son obligatorios'}, status=400)

            if not email.endswith('@ejercito.mil.co'):
                return JsonResponse({'error': 'Correo institucional requerido'}, status=400)

            if User.objects.filter(username=email).exists():
                return JsonResponse({'error': 'El usuario ya existe'}, status=400)

            if Perfil.objects.filter(documento=documento).exists():
                return JsonResponse({'error': 'Documento ya registrado'}, status=400)

            
            user = User.objects.create(
                username=email,
                email=email,
                first_name=nombres,
                last_name=apellidos,
                password=make_password(documento) 
            )

            Perfil.objects.create(
                user=user,
                rol=rol,
                documento=documento,
                unidad=unidad,
                grado=grado
            )

            return JsonResponse({'mensaje': 'Usuario creado correctamente'})

        except Exception as e:
            return JsonResponse({'error': str(e)}, status=500)
        
@csrf_exempt
def cargar_usuarios_excel(request):
    if request.method == "POST":

        archivo = request.FILES.get("archivo")
        rol = request.POST.get("rol")

        if not archivo:
            return JsonResponse({
                "error": "No se recibió ningún archivo"
            })

        if not rol:
            return JsonResponse({
                "error": "Debe seleccionar un tipo de usuario"
            })

        try: 
            df = pd.read_excel(archivo, header=7) 

            creados = 0
            errores = []

            for index, row in df.iterrows():

                nombres = str(row.get("nombres", "")).strip()
                apellidos = str(row.get("apellidos", "")).strip()
                email = str(row.get("email", "")).strip().lower()
                documento = str(row.get("documento", "")).strip()
                grado = str(row.get("grado", "")).strip()
                unidad = str(row.get("unidad", "")).strip()
                batallon = str(row.get("batallon", "")).strip()
                compania = str(row.get("compania", "")).strip()

                if not unidad:
                    unidad = " - ".join([valor for valor in [batallon, compania] if valor])


                print( f"Fila {index+2}:", 
                      nombres, 
                      apellidos, 
                      email, 
                      documento, 
                      grado, 
                      unidad 
                    )
                
                print("GRADO:", grado)

                if not email:
                    errores.append(f"Fila {index+2}: email vacío")
                    continue

                if not documento:
                    errores.append(f"Fila {index+2}: documento vacío")
                    continue

                if not email.endswith("@ejercito.mil.co"):
                    errores.append(f"Fila {index+2}: correo no institucional")
                    continue

                if User.objects.filter(username=email).exists():
                    errores.append(f"Fila {index+2}: usuario ya existe")
                    continue

                if Perfil.objects.filter(documento=documento).exists():
                    errores.append(f"Fila {index+2}: documento ya registrado")
                    continue

                user = User.objects.create(
                    username=email,
                    email=email,
                    first_name=nombres,
                    last_name=apellidos,
                    password=make_password(documento)
                )

                Perfil.objects.create(
                    user=user,
                    rol=rol,
                    documento=documento,
                    grado=grado,
                    unidad=unidad
                )

                creados += 1

            return JsonResponse({
                "mensaje": f"{creados} usuarios creados",
                "errores": errores
            })

        except Exception as e:
            return JsonResponse({
                "error": str(e)
            }, status=500)

    return JsonResponse({
        "error": "Método no permitido"
    })

@csrf_exempt
def descargar_plantilla_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Plantilla Usuarios"

    ruta_logo_principal = os.path.join(
        settings.BASE_DIR,
        "static",
        "img",
        "escudo.png"
    )

    if os.path.exists(ruta_logo_principal):
        logo_principal = Image(ruta_logo_principal)
        logo_principal.width = 140
        logo_principal.height = 140

        # Logo izquierdo superior
        ws.add_image(logo_principal, "A1")


    ruta_logo_secundario = os.path.join(
        settings.BASE_DIR,
        "static",
        "img",
        "CompromisoColombia.png"
    )

    if os.path.exists(ruta_logo_secundario):
        logo_secundario = Image(ruta_logo_secundario)
        logo_secundario.width = 180
        logo_secundario.height = 110

        # Logo derecho superior
        ws.add_image(logo_secundario, "E2")

    encabezados = [
        "nombres",
        "apellidos",
        "email",
        "documento",
        "grado",
        "batallon",
        "compania"
    ]

    for i, titulo in enumerate(encabezados, start=1):
        ws.cell(row=8, column=i, value=titulo)


    ejemplo = [
        "Juan",
        "Perez",
        "juan@ejercito.mil.co",
        "12345678",
        "Capitán",
        "Batallón Norte",
        "Compañía A"
    ]

    ws["A9"] = "Juan"
    ws["B9"] = "Perez"
    ws["C9"] = "juan@ejercito.mil.co"
    ws["D9"] = "12345678"
    ws["E9"] = "Capitán"
    ws["F9"] = "Batallón Norte"
    ws["G9"] = "Compañía A"

    for i, valor in enumerate(ejemplo, start=1):
        ws.cell(row=9, column=i, value=valor)

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="789441"
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
        size=12
    )

    center = Alignment(
        horizontal="center",
        vertical="center"
    )

    thin = Side(
        style="thin",
        color="D9D9D9"
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin
    )

    for col in ws[8]:
        col.fill = header_fill
        col.font = header_font
        col.alignment = center
        col.border = border


    for row in ws.iter_rows(min_row=9, max_row=9):
        for cell in row:
            cell.border = border
            cell.alignment = center

    widths = {
        "A": 22,
        "B": 22,
        "C": 35,
        "D": 20,
        "E": 28,
        "F": 28,
        "G": 28
    }

    for col, width in widths.items():
        ws.column_dimensions[col].width = width


    for fila in range(10, 501):
        for col in ["A", "B", "C", "D", "E", "F", "G"]:
            ws[f"{col}{fila}"] = ""

    ws.protection.sheet = True
    ws.protection.enable()

    # desbloquear celdas editables
    for fila in range(9, 501):
        for col in ["A", "B", "C", "D", "E", "F", "G"]:
            ws[f"{col}{fila}"].protection = (
                ws[f"{col}{fila}"].protection.copy(locked=False)
            )

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response["Content-Disposition"] = (
        'attachment; filename="plantilla_usuarios.xlsx"'
    )

    wb.save(response)

    return response


def obtener_estadisticas(request):
    total_usuarios = User.objects.count()
    total_soldados = Perfil.objects.filter(rol="soldado").count()
    total_instructores = Perfil.objects.filter(rol="instructor").count()
    total_activos = User.objects.filter(is_active=True).count()

    return JsonResponse({
        "total_usuarios": total_usuarios,
        "total_soldados": total_soldados,
        "total_instructores": total_instructores,
        "total_activos": total_activos
    })

def listar_usuarios(request):
    usuarios = Perfil.objects.select_related("user").all()

    return render(
        request,
        "admin_usuarios.html",
        {
            "usuarios": usuarios
        }
    )

@csrf_exempt
def eliminar_usuario(request, user_id):
    try:
        user = User.objects.get(id=user_id)
        user.delete()

        return JsonResponse({
            "mensaje": "Usuario eliminado correctamente"
        })

    except User.DoesNotExist:
        return JsonResponse({
            "error": "Usuario no encontrado"
        }, status=404)

@csrf_exempt
def editar_usuario(request, user_id):
    if request.method == "POST":
        try:
            user = User.objects.get(id=user_id)
            perfil = Perfil.objects.get(user=user)

            data = json.loads(request.body)

            user.first_name = data.get("nombre")
            user.last_name = data.get("apellido")
            user.email = data.get("email")
            user.is_active = data.get("activo")

            perfil.rol = data.get("rol")
            perfil.documento = data.get("documento") or perfil.documento
            perfil.grado = data.get("grado")
            perfil.unidad = data.get("unidad")

            user.save()
            perfil.save()

            return JsonResponse({
                "mensaje": "Usuario actualizado correctamente"
            })

        except User.DoesNotExist:
            return JsonResponse({
                "error": "Usuario no encontrado"
            }, status=404)

@csrf_exempt
def crear_curso(request):

    if request.method != "POST":
        return JsonResponse({
            "error": "Método no permitido"
        }, status=405)

    try:
        data = json.loads(request.body)

        codigo = data.get("codigo")
        nombre = data.get("nombre")
        descripcion = data.get("descripcion")
        instructor_id = data.get("instructor")
        fecha_inicio = data.get("fecha_inicio")
        fecha_fin = data.get("fecha_fin")
        cupo_maximo = data.get("cupo_maximo")

        # Validaciones
        if not codigo or not nombre:
            return JsonResponse({
                "error": "Código y nombre son obligatorios"
            }, status=400)

        # Curso duplicado
        if Curso.objects.filter(codigo=codigo).exists():
            return JsonResponse({
                "error": "Ya existe un curso con ese código"
            }, status=400)

        # Instructor válido
        try:
            instructor = User.objects.get(id=instructor_id)

        except User.DoesNotExist:
            return JsonResponse({
                "error": "Instructor no encontrado"
            }, status=404)

        # Crear curso
        Curso.objects.create(
            codigo=codigo,
            nombre=nombre,
            descripcion=descripcion,
            instructor=instructor,
            fecha_inicio=fecha_inicio,
            fecha_fin=fecha_fin,
            cupo_maximo=cupo_maximo
        )

        return JsonResponse({
            "mensaje": "Curso creado correctamente"
        }, status=201)

    except json.JSONDecodeError:
        return JsonResponse({
            "error": "JSON inválido"
        }, status=400)

    except Exception as e:
        return JsonResponse({
            "error": str(e)
        }, status=500)

def listar_cursos(request):

    cursos = Curso.objects.select_related('instructor').all()

    data = []

    for curso in cursos:
        data.append({
            "id": curso.id,
            "codigo": curso.codigo,
            "nombre": curso.nombre,
            "instructor": curso.instructor.first_name if curso.instructor else "Sin instructor",
            "fecha_inicio": curso.fecha_inicio,
            "fecha_fin": curso.fecha_fin,
            "estado": "Activo" if curso.activo else "Inactivo"
        })

    return JsonResponse({
        "data": data
    })

@csrf_exempt
def editar_curso(request, curso_id):

    if request.method != "POST":
        return JsonResponse({
            "error": "Método no permitido"
        }, status=405)

    try:

        curso = Curso.objects.get(id=curso_id)

        data = json.loads(request.body)

        curso.codigo = data.get("codigo")
        curso.nombre = data.get("nombre")
        curso.descripcion = data.get("descripcion")
        curso.fecha_inicio = data.get("fecha_inicio")
        curso.fecha_fin = data.get("fecha_fin")
        curso.cupo_maximo = data.get("cupo_maximo")

        instructor_id = data.get("instructor")

        curso.instructor = User.objects.get(
            id=instructor_id
        )

        curso.save()

        return JsonResponse({
            "mensaje": "Curso actualizado correctamente"
        })

    except Curso.DoesNotExist:
        return JsonResponse({
            "error": "Curso no encontrado"
        }, status=404)

    except Exception as e:
        return JsonResponse({
            "error": str(e)
        }, status=500)
    
@csrf_exempt
def eliminar_curso(request, curso_id):

    try:

        curso = Curso.objects.get(id=curso_id)

        curso.delete()

        return JsonResponse({
            "mensaje": "Curso eliminado correctamente"
        })

    except Curso.DoesNotExist:

        return JsonResponse({
            "error": "Curso no encontrado"
        }, status=404)



